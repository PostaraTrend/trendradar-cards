import io
import os
import tempfile
import unittest
from pathlib import Path

from flask import Flask
from docx import Document
import fitz
import resume_export as renderer

app = Flask(__name__)
app.register_blueprint(renderer.export_bp)
CLIENT = app.test_client()
BASE = {
    'seeker_name': 'Alex Example',
    'contact_lines': ['alex@example.test', 'Edmonton, Alberta'],
    'posting': {'title': 'Operations Coordinator', 'employer': 'Example Company', 'location': 'Edmonton'},
    'sections': [{'title': 'Experience', 'lines': [
        {'text': 'Coordinated schedules for a team of twelve colleagues.', 'heading': False},
        {'text': 'Maintained accurate records and prepared weekly reports.', 'heading': False}]}],
    'body': ['Dear Hiring Manager,', 'I coordinated schedules for a team of twelve colleagues.',
             'I maintained accurate records and prepared weekly reports.', 'Thank you for considering my application.'],
    'date': 'October 2, 2026',
}


class DocumentStyles(unittest.TestCase):
    def test_outputs_and_matching_word_layout(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        out = Path(os.environ.get('TS_STYLE_SAMPLE_DIR', temporary.name))
        out.mkdir(exist_ok=True)
        for style in ('classic', 'compact', 'plain'):
            word_docs = []
            for kind in ('resume', 'cover'):
                for fmt in ('pdf', 'docx'):
                    with self.subTest(style=style, kind=kind, format=fmt):
                        response = CLIENT.post('/export/' + kind, json={**BASE, 'template': style, 'format': fmt})
                        self.assertEqual(response.status_code, 200)
                        (out / f'{kind}-{style}.{fmt}').write_bytes(response.data)
                        if fmt == 'pdf':
                            pdf = fitz.open(stream=response.data, filetype='pdf')
                            text = ''.join(page.get_text() for page in pdf)
                            self.assertEqual(len(pdf), 1)
                            fonts = ' '.join(f[3] for page in pdf for f in page.get_fonts())
                            self.assertIn('Times' if style == 'plain' else 'Helvetica', fonts)
                            pdf.close()
                        else:
                            doc = Document(io.BytesIO(response.data))
                            word_docs.append(doc)
                            text = '\n'.join(p.text for p in doc.paragraphs)
                            self.assertEqual(doc.styles['Normal'].font.name, renderer._TEMPLATES[style]['docx_font'])
                            self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, renderer._TEMPLATES[style]['name'])
                        self.assertIn(BASE['seeker_name'], text)
                        self.assertIn(BASE['contact_lines'][0], text)
                        lines = [x['text'] for x in BASE['sections'][0]['lines']] if kind == 'resume' else BASE['body']
                        normalized = ' '.join(text.split())
                        for line in lines:
                            self.assertIn(line, normalized)
                        self.assertNotIn('Verified against', text)
            a, b = [d.sections[0] for d in word_docs]
            for attr in ('page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
                self.assertEqual(getattr(a, attr), getattr(b, attr))

    def test_default_and_invalid_template(self):
        for kind in ('resume', 'cover'):
            self.assertEqual(CLIENT.post('/export/' + kind, json=BASE).status_code, 200)
            self.assertEqual(CLIENT.post('/export/' + kind, json={**BASE, 'template': 'invented'}).status_code, 422)

    def test_verification_footer_is_still_gated(self):
        for kind in ('resume', 'cover'):
            for metadata in ({'verified_at': 'October 2, 2026'}, {'attested_version': 3},
                             {'verified_at': 'October 2, 2026', 'attested_version': 3}):
                r = CLIENT.post('/export/' + kind, json={**BASE, **metadata, 'format': 'docx'})
                self.assertEqual(r.status_code, 200)
                text = '\n'.join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
                self.assertEqual('verified' in text.lower(), len(metadata) == 2)

    def test_long_cover_retains_every_sentence(self):
        lines = ['Dear Hiring Manager,'] + [f'Example responsibility {i}: prepared accurate weekly records for the operations team.' for i in range(100)] + ['Thank you.']
        for style in renderer._TEMPLATES:
            r = CLIENT.post('/export/cover', json={**BASE, 'body': lines, 'template': style})
            self.assertEqual(r.status_code, 200)
            pdf = fitz.open(stream=r.data, filetype='pdf')
            text = ' '.join(''.join(page.get_text() for page in pdf).split())
            self.assertGreater(len(pdf), 1)
            for line in lines:
                self.assertIn(line, text)


if __name__ == '__main__':
    unittest.main(verbosity=2)
